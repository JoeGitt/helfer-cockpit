import io
import json
import threading
import urllib.error
import urllib.request
from pathlib import Path
import openpyxl
import pytest
from cockpit.webapp import Zustand, baue_dashboard, starte_server

FG_KOPF = ["Kontakte", "Kontakt-ID Verein", "Primäre E-Mail", "Vorname", "Nachname",
          "Handy", "Mitgliedschaft", "E-Mail Eltern 1", "E-Mail Eltern 2", "Geburtsdatum"]


def _fairgate_bytes(zeilen):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(FG_KOPF)
    for z in zeilen:
        ws.append(z)
    puffer = io.BytesIO()
    wb.save(puffer)
    return puffer.getvalue()

FIXTURES = Path(__file__).parent / "fixtures"

def _zustand(tmp_path):
    return Zustand(
        helpers=json.loads((FIXTURES / "helpers.json").read_text(encoding="utf-8")),
        assignments=json.loads((FIXTURES / "assignments.json").read_text(encoding="utf-8")),
        stand="2026-08-31 14:32",
        regeln_pfad=tmp_path / "regeln.json",
        protokoll_pfad=tmp_path / "protokoll.jsonl",
        ausgabe_dir=tmp_path / "Ausgabe",
        api_client_factory=None,
        fehler="")

def test_dashboard_json(tmp_path):
    d = baue_dashboard(_zustand(tmp_path))
    assert d["kennzahlen"]["mitglieder"] == 3
    fg2417 = next(m for m in d["mitglieder"] if m["fg"] == "FG-2417")
    assert fg2417["ist"] == 2.0 and fg2417["status_saison"] == "erfuellt"
    assert len(fg2417["accounts"]) == 3
    assert d["wer_leistet"]["zweitaccount"] == 2      # num_ok von 101
    assert any(h["code"] == "D6" for h in d["hinweise"])
    assert d["kennzahlen"]["nok_summe"] == 1           # Bericht-Kachel NOK (Spez. 6.2)

def test_dashboard_ohne_abruf(tmp_path):
    z = _zustand(tmp_path)
    z.helpers = None
    d = baue_dashboard(z)
    assert d["stand"] == "" or d["mitglieder"] == []

@pytest.fixture
def server(tmp_path):
    srv = starte_server(_zustand(tmp_path), port=0)
    t = threading.Thread(target=srv.serve_forever, daemon=True)
    t.start()
    yield f"http://127.0.0.1:{srv.server_address[1]}"
    srv.shutdown()

def test_api_stand_http(server):
    with urllib.request.urlopen(server + "/api/stand") as r:
        d = json.loads(r.read())
    assert d["kennzahlen"]["mitglieder"] == 3

def test_index_wird_ausgeliefert(server):
    with urllib.request.urlopen(server + "/") as r:
        html = r.read().decode("utf-8")
    assert "HELFER-COCKPIT" in html and "app.js" in html

def test_bindet_nur_lokal(tmp_path):
    srv = starte_server(_zustand(tmp_path), port=0)
    assert srv.server_address[0] == "127.0.0.1"
    srv.server_close()


class _LeererClient:
    """Stub: liefert 0 Helfende (Spez. 6.4: unplausibel, Abbruch)."""
    def helpers(self):
        return []


def test_gesamtexport_http(server):
    req = urllib.request.Request(server + "/api/export/gesamt", method="POST", data=b"")
    with urllib.request.urlopen(req) as r:
        d = json.loads(r.read())
    assert d["anzahl"] == 3
    assert Path(d["datei"]).is_file()
    assert Path(d["datei"]).name.startswith("gesamtexport-")


def test_fairgate_upload_liefert_unbekannte_kategorien_feld(server):
    daten = _fairgate_bytes([
        ["x", 9999, "neu@example.ch", "Fremd", "Kategorie", None,
         "Vereinsfremd", None, None, "2000-01-01"],
    ])
    req = urllib.request.Request(server + "/api/fairgate", method="POST", data=daten)
    with urllib.request.urlopen(req) as r:
        d = json.loads(r.read())
    assert d["unbekannte_kategorien"] and "Vereinsfremd" in d["unbekannte_kategorien"][0]
    assert d["neueintritte"] == 0


def test_fairgate_garbage_bytes_liefert_400(server):
    req = urllib.request.Request(server + "/api/fairgate", method="POST",
                                 data=b"nicht mal ansatzweise eine Excel-Datei \x00\x01")
    try:
        urllib.request.urlopen(req)
        assert False, "erwartete HTTPError 400"
    except urllib.error.HTTPError as e:
        assert e.code == 400
        d = json.loads(e.read())
        assert d["fehler"]


class _EventsWirftClient:
    """Stub: helpers() liefert normal, events() scheitert (D6 degradiert, I3)."""
    def __init__(self, helpers):
        self._helpers = helpers

    def helpers(self):
        return self._helpers

    def events(self):
        raise RuntimeError("API für Einsätze nicht erreichbar")


def test_d6_hinweis_bei_assignments_fehler(tmp_path):
    z = _zustand(tmp_path)
    z.api_client_factory = lambda: _EventsWirftClient(z.helpers)
    srv = starte_server(z, port=0)
    t = threading.Thread(target=srv.serve_forever, daemon=True)
    t.start()
    try:
        basis = f"http://127.0.0.1:{srv.server_address[1]}"
        req = urllib.request.Request(basis + "/api/abruf", method="POST", data=b"")
        with urllib.request.urlopen(req) as r:
            d = json.loads(r.read())
        assert not d["fehler"]      # helpers() gelang, nur assignments fehlten
        d6 = [h for h in d["hinweise"] if h["code"] == "D6" and "nicht möglich" in h["text"]]
        assert len(d6) == 1
    finally:
        srv.shutdown()


class _OkClient:
    def __init__(self, helpers, events=None):
        self._helpers, self._events = helpers, (events or [])

    def helpers(self):
        return self._helpers

    def events(self):
        return self._events

    def alle_assignments(self, events):
        return []


def test_d6_degradations_hinweis_verschwindet_nach_erfolgreichem_abruf(tmp_path):
    z = _zustand(tmp_path)
    z.assignments_fehler = True     # simuliert einen vorherigen fehlgeschlagenen Abruf
    z.api_client_factory = lambda: _OkClient(z.helpers)
    srv = starte_server(z, port=0)
    t = threading.Thread(target=srv.serve_forever, daemon=True)
    t.start()
    try:
        basis = f"http://127.0.0.1:{srv.server_address[1]}"
        req = urllib.request.Request(basis + "/api/abruf", method="POST", data=b"")
        with urllib.request.urlopen(req) as r:
            d = json.loads(r.read())
        assert not any(h["code"] == "D6" and "nicht möglich" in h["text"] for h in d["hinweise"])
    finally:
        srv.shutdown()


def test_korruptes_regeln_json_liefert_standardregeln_und_hinweis(tmp_path):
    z = _zustand(tmp_path)
    z.regeln_pfad.parent.mkdir(parents=True, exist_ok=True)
    z.regeln_pfad.write_text("{das ist kein json", encoding="utf-8")
    srv = starte_server(z, port=0)
    t = threading.Thread(target=srv.serve_forever, daemon=True)
    t.start()
    try:
        basis = f"http://127.0.0.1:{srv.server_address[1]}"
        with urllib.request.urlopen(basis + "/api/stand") as r:
            d = json.loads(r.read())
        assert d["kennzahlen"]["mitglieder"] == 3   # trotzdem geladen
        assert any("Regeln" in h["text"] or "beschädigt" in h["text"] for h in d["hinweise"])
        with urllib.request.urlopen(basis + "/api/regeln") as r:
            regeln = json.loads(r.read())
        assert regeln["altersgrenze"] == 16          # Standardregeln
    finally:
        srv.shutdown()


def test_kaputte_route_liefert_500_json_statt_abgerissener_verbindung(server):
    req = urllib.request.Request(server + "/api/regeln", method="POST",
                                 data=b"das ist kein json", headers={"Content-Type": "application/json"})
    try:
        urllib.request.urlopen(req)
        assert False, "erwartete HTTPError 500"
    except urllib.error.HTTPError as e:
        assert e.code == 500
        d = json.loads(e.read())
        assert d["fehler"]


def test_fremder_host_header_wird_abgewiesen(server):
    req = urllib.request.Request(server + "/api/stand", headers={"Host": "evil.example.com"})
    try:
        urllib.request.urlopen(req)
        assert False, "erwartete HTTPError 403"
    except urllib.error.HTTPError as e:
        assert e.code == 403


def test_normaler_request_kommt_trotz_host_guard_durch(server):
    with urllib.request.urlopen(server + "/api/stand") as r:
        assert r.status == 200


def test_ungueltige_sicht_liefert_400_statt_keyerror(server):
    req = urllib.request.Request(server + "/api/export/saeumige?sicht=galaxie", method="POST", data=b"")
    try:
        urllib.request.urlopen(req)
        assert False, "erwartete HTTPError 400"
    except urllib.error.HTTPError as e:
        assert e.code == 400
        d = json.loads(e.read())
        assert d["fehler"]


def test_api_abruf_log_zaehlt_hinweise(tmp_path):
    z = _zustand(tmp_path)
    z.api_client_factory = lambda: _OkClient(z.helpers)
    srv = starte_server(z, port=0)
    t = threading.Thread(target=srv.serve_forever, daemon=True)
    t.start()
    try:
        basis = f"http://127.0.0.1:{srv.server_address[1]}"
        req = urllib.request.Request(basis + "/api/abruf", method="POST", data=b"")
        with urllib.request.urlopen(req) as r:
            d = json.loads(r.read())
        with urllib.request.urlopen(basis + "/api/protokoll") as r:
            protokoll_liste = json.loads(r.read())
        eintrag = next(e for e in protokoll_liste if e["aktion"] == "api-abruf")
        assert eintrag["hinweise"] == d["kennzahlen"]["hinweise"]
    finally:
        srv.shutdown()


def test_abruf_fehler_lasst_alte_anzeige_stehen(tmp_path):
    z = _zustand(tmp_path)
    z.api_client_factory = lambda: _LeererClient()
    srv = starte_server(z, port=0)
    t = threading.Thread(target=srv.serve_forever, daemon=True)
    t.start()
    try:
        basis = f"http://127.0.0.1:{srv.server_address[1]}"
        req = urllib.request.Request(basis + "/api/abruf", method="POST", data=b"")
        with urllib.request.urlopen(req) as r:
            d = json.loads(r.read())
        assert d["fehler"]
        assert d["kennzahlen"]["mitglieder"] == 3
    finally:
        srv.shutdown()


# ---------------------------------------------------------------------------
# Erweiterung «top notch»: alle Accounts, Portal-Links, Kategorie-Erfüllung,
# Ausgabe-Dateien im Browser öffnen, drittes Export-Blatt.
# ---------------------------------------------------------------------------

def test_stand_enthaelt_alle_accounts_mit_portal_link(tmp_path):
    z = _zustand(tmp_path)
    z.org_slug = "pfadi-winterthur-handball"
    d = baue_dashboard(z)
    assert len(d["alle_accounts"]) == 10
    a = next(x for x in d["alle_accounts"] if x["id"] == 101)
    assert a["typ"] == "zweitaccount" and a["fg"] == "FG-2417" and a["num_ok"] == 2
    assert a["email"] == "reto@example.ch"
    assert a["portal_url"] == "https://app.helfereinsatz.ch/pfadi-winterthur-handball/de/helpers/detail/101"
    # auch die Accounts innerhalb der Mitglieder tragen id + Portal-Link
    m = next(x for x in d["mitglieder"] if x["fg"] == "FG-2417")
    assert all("portal_url" in acc and "id" in acc for acc in m["accounts"])


def _fairgate_xlsx_bytes(zeilen):
    import io
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Kontakte", "Kontakt-ID Verein", "Primäre E-Mail", "Vorname", "Nachname",
               "Handy", "Mitgliedschaft", "E-Mail Eltern 1", "E-Mail Eltern 2", "Geburtsdatum"])
    for z in zeilen:
        ws.append(z)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_kategorie_erfuellung_nur_mit_fairgate_export(server):
    with urllib.request.urlopen(server + "/api/stand") as r:
        assert json.loads(r.read())["kategorie_erfuellung"] == []
    daten = _fairgate_xlsx_bytes([
        ["x", 2417, "lina@example.ch", "Lina", "Brunner", "", "Junioren", "eltern@example.ch", None, "2000-01-01"],
        ["x", 1083, "noah@example.ch", "Noah", "Keller", "", "Aktivmitglied", None, None, "2000-01-01"],
    ])
    req = urllib.request.Request(server + "/api/fairgate", data=daten, method="POST")
    with urllib.request.urlopen(req) as r:
        assert r.status == 200
    with urllib.request.urlopen(server + "/api/stand") as r:
        ke = json.loads(r.read())["kategorie_erfuellung"]
    assert {"kategorie": "Junioren", "gesamt": 1, "erreicht": 1} in ke
    assert {"kategorie": "Aktivmitglied", "gesamt": 1, "erreicht": 1} in ke
    assert not any(k["kategorie"] == "" for k in ke)


def test_ausgabe_datei_wird_ausgeliefert_und_traversal_blockiert(server, tmp_path):
    ausgabe = tmp_path / "Ausgabe"
    ausgabe.mkdir(exist_ok=True)
    (ausgabe / "probe.html").write_text("<p>Probe</p>", encoding="utf-8")
    (tmp_path / "geheim.txt").write_text("nein", encoding="utf-8")
    with urllib.request.urlopen(server + "/ausgabe/probe.html") as r:
        assert r.status == 200 and b"Probe" in r.read()
        assert "text/html" in r.headers.get("Content-Type", "")
    import http.client
    host, port = server.replace("http://", "").split(":")
    c = http.client.HTTPConnection(host, int(port))
    c.request("GET", "/ausgabe/../geheim.txt")
    assert c.getresponse().status == 404
    c.request("GET", "/ausgabe/gibtsnicht.html")
    assert c.getresponse().status == 404


def test_gesamtexport_hat_blatt_alle_helfenden(server):
    import openpyxl
    req = urllib.request.Request(server + "/api/export/gesamt", data=b"", method="POST")
    with urllib.request.urlopen(req) as r:
        pfad = json.loads(r.read())["datei"]
    wb = openpyxl.load_workbook(pfad)
    assert wb.sheetnames == ["Mitglieder", "Accounts", "Alle Helfenden"]
    ws = wb["Alle Helfenden"]
    assert ws.max_row == 11  # Kopf + 10 Accounts
    assert ws.cell(1, 1).value == "ID" and "Typ" in [c.value for c in ws[1]]


def test_stand_meldet_api_verfuegbarkeit_und_letzten_abgleich(tmp_path):
    z = _zustand(tmp_path)
    d = baue_dashboard(z)
    assert d["api_verfuegbar"] is False and d["letzter_abgleich"] is None


def test_fairgate_liefert_kategorien_abweichungen_und_kontrolle(server):
    daten = _fairgate_xlsx_bytes([
        ["x", 2417, "lina.neu@example.ch", "Lina", "Brunner", "", "Aktivmitglied", None, None, "2000-01-01"],
        ["x", 1083, "noah@example.ch", "Noah", "Keller", "", "Junioren", None, None, "2000-01-01"],
        ["x", 3105, "sara@example.ch", "Sara", "Meili", "", "Aktivmitglied", None, None, "2000-01-01"],
        ["x", 8888, "x@example.ch", "Gast", "Gönner", "", "Gönner", None, None, "2000-01-01"],
    ])
    req = urllib.request.Request(server + "/api/fairgate", data=daten, method="POST")
    with urllib.request.urlopen(req) as r:
        d = json.loads(r.read())
    assert d["kategorien"]["pflichtig"] == {"Aktivmitglied": 2, "Junioren": 1}
    assert d["kategorien"]["unbekannt"] == {"Gönner": 1}
    assert d["kontakt_abweichungen"][0]["fg"] == "FG-2417"           # Lina: Portal lina@, Fairgate lina.neu@
    assert all("portal_url" in h and "helper_id" in h for h in d["handarbeit"])
    assert d["dateien"]["kontakte"].endswith(".csv")
    with urllib.request.urlopen(server + "/api/stand") as r:
        st = json.loads(r.read())
    assert st["letzter_abgleich"]["zusammenfassung"] == d["zusammenfassung"]
    # Kontrolle: gleiche Kontakte gegen den (unveränderten) Portal-Bestand → nicht synchron,
    # weil Korrekturen (Zweitaccount/Freiwillige mit Zielwert) weiterhin anstehen
    req = urllib.request.Request(server + "/api/abgleich/kontrolle", data=b"", method="POST")
    with urllib.request.urlopen(req) as r:
        k = json.loads(r.read())
    assert k["synchron"] is False and k["offen"]["korrekturen"] >= 1
    assert "zusammenfassung" in k


def test_kontrolle_ohne_fairgate_liefert_400(server):
    import http.client
    host, port = server.replace("http://", "").split(":")
    c = http.client.HTTPConnection(host, int(port))
    c.request("POST", "/api/abgleich/kontrolle", body=b"")
    assert c.getresponse().status == 400


def test_regeln_post_mit_email_abweichung(server):
    body = json.dumps({"kategorien": [], "altersgrenze": 16, "halbjahresziel": 1,
                       "email_abweichung": "handarbeit"}).encode()
    req = urllib.request.Request(server + "/api/regeln", data=body, method="POST",
                                 headers={"Content-Type": "application/json"})
    with urllib.request.urlopen(req) as r:
        assert json.loads(r.read())["ok"] is True
    with urllib.request.urlopen(server + "/api/regeln") as r:
        assert json.loads(r.read())["email_abweichung"] == "handarbeit"


def test_letzter_abgleich_ist_nach_upload_abrufbar(server):
    import http.client
    host, port = server.replace("http://", "").split(":")
    c = http.client.HTTPConnection(host, int(port))
    c.request("GET", "/api/abgleich/letzter")
    assert c.getresponse().status == 404                      # noch kein Abgleich in dieser Sitzung
    daten = _fairgate_xlsx_bytes([
        ["x", 2417, "lina@example.ch", "Lina", "Brunner", "", "Aktivmitglied", None, None, "2000-01-01"],
    ])
    req = urllib.request.Request(server + "/api/fairgate", data=daten, method="POST")
    with urllib.request.urlopen(req) as r:
        d = json.loads(r.read())
    with urllib.request.urlopen(server + "/api/abgleich/letzter") as r:
        l = json.loads(r.read())
    assert l["zusammenfassung"] == d["zusammenfassung"] and l["dateien"] == d["dateien"]
    assert "handarbeit" in l and "kategorien" in l


def test_fairgate_liefert_import_vorschau(server):
    daten = _fairgate_xlsx_bytes([
        ["x", 9999, "neu@example.ch", "Neu", "Kind", "0790000000", "Aktivmitglied", None, None, "2000-01-01"],
    ])
    req = urllib.request.Request(server + "/api/fairgate", data=daten, method="POST")
    with urllib.request.urlopen(req) as r:
        d = json.loads(r.read())
    v = d["import_vorschau"]
    neu = next(z for z in v if z["art"] == "Neueintritt")
    assert neu["name"] == "Neu Kind" and "FG-9999" in neu["aenderungen"] and "Zielwert 2" in neu["aenderungen"]
    assert any(z["art"] == "Korrektur" for z in v)          # Zweitaccount/Freiwillige mit Zielwert → 0
    assert [z["zeile"] for z in v] == list(range(2, len(v) + 2))       # Excel-Zeilen, Kopfzeile = 1
    assert neu["kategorie"] == "Neueintritt" and all(z["kategorie"] and "email" in z for z in v)
    # Begleitdatei mit Begründung liegt neben der Import-Datei und wird ausgeliefert
    assert d["dateien"]["begruendung"].endswith("-begruendung.html")
    with urllib.request.urlopen(server + "/ausgabe/" + d["dateien"]["begruendung"].split("/")[-1]) as r:
        html = r.read().decode()
    assert "Neu Kind" in html and "Zeile" in html and "nicht ins Portal importiert" in html


# ---- Vorfragen (möglicher Zweitaccount) vor der Import-Datei beantworten ----------------

def _acc_json(id, vn, nn, mail, fg, gruppen=("Mitglied",), ziel=2):
    return {"id": id, "firstName": vn, "lastName": nn, "email": mail, "adminRemarks": fg or "",
            "groups": [{"id": i, "name": g} for i, g in enumerate(gruppen)],
            "stateCache": {"requestedValue": ziel, "plannedValue": 0}}

def _server_mit(tmp_path, helpers):
    z = _zustand(tmp_path)
    z.helpers = helpers
    z.assignments = []
    srv = starte_server(z, port=0)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    return srv, f"http://127.0.0.1:{srv.server_address[1]}"

def _post_json(url, daten):
    req = urllib.request.Request(url, data=json.dumps(daten).encode(), method="POST",
                                 headers={"Content-Type": "application/json"})
    with urllib.request.urlopen(req) as r:
        return json.loads(r.read())

def test_vorfragen_endpoint_erzeugt_import_erst_nach_antwort(tmp_path):
    helpers = [_acc_json(1, "Lina", "Brunner", "lina@example.ch", "FG-1"),
               _acc_json(2, "Reto", "Brunner", "reto@example.ch", None),
               _acc_json(3, "Noah", "Keller", "noah@example.ch", "FG-2")]
    fairgate = _fairgate_xlsx_bytes([
        ["x", 1, "lina@example.ch", "Lina", "Brunner", "079 1", "Aktivmitglied", None, None, "2000-01-01"],
        ["x", 2, "noah@example.ch", "Noah", "Keller", "", "Aktivmitglied", None, None, "2000-01-01"]])
    srv, url = _server_mit(tmp_path, helpers)
    try:
        # ohne Fairgate: 400
        try:
            _post_json(url + "/api/abgleich/entscheide", {"entscheide": {}})
            assert False
        except urllib.error.HTTPError as e:
            assert e.code == 400
        req = urllib.request.Request(url + "/api/fairgate", data=fairgate, method="POST")
        with urllib.request.urlopen(req) as r:
            d = json.loads(r.read())
        assert len(d["vorfragen"]) == 1
        v = d["vorfragen"][0]
        assert v["helper_id"] == 2 and v["portal_url"].endswith("/detail/2")
        assert v["kandidaten"][0]["fg"] == "FG-1" and v["kandidaten"][0]["telefon"] == "079 1"
        assert v["kandidaten"][0]["portal_url"].endswith("/detail/1")
        assert not any(z["name"] == "Reto Brunner" for z in d["import_vorschau"])
        # Antwort Zweitaccount → Import-Zeile mit FG, Datei neu geschrieben
        d = _post_json(url + "/api/abgleich/entscheide", {"entscheide": {"2": {"antwort": "zweitaccount", "fg": "FG-1", "name": "Reto Brunner"}}})
        assert d["vorfragen"] == []
        z = next(z for z in d["import_vorschau"] if z["name"] == "Reto Brunner")
        assert "Bemerkung FG-1" in z["aenderungen"] and z["kategorie"] == "Zweitaccount"
        assert d["entscheide"]["2"]["antwort"] == "zweitaccount" and d["entscheide"]["2"]["gespeichert"] is False
        wb = openpyxl.load_workbook(d["dateien"]["import"])
        zeilen = list(wb.active.iter_rows(values_only=True))
        assert any(r[0] == "Reto" and r[10] == "FG-1" for r in zeilen)
        # Antwort ändern auf «andere» → dauerhaft gespeichert
        d = _post_json(url + "/api/abgleich/entscheide", {"entscheide": {"2": {"antwort": "andere", "name": "Reto Brunner"}}})
        z = next(z for z in d["import_vorschau"] if z["name"] == "Reto Brunner")
        assert "Bemerkung" not in z["aenderungen"] and d["entscheide"]["2"]["gespeichert"] is True
        gespeichert = json.loads((tmp_path / "Ausgabe" / "entscheide.json").read_text())
        assert gespeichert["2"]["antwort"] == "andere" and gespeichert["2"]["name"] == "Reto Brunner"
        # Kontrolle kennt die Antwort ebenfalls (keine Vorfrage offen)
        req = urllib.request.Request(url + "/api/abgleich/kontrolle", data=b"{}", method="POST")
        with urllib.request.urlopen(req) as r:
            k = json.loads(r.read())
        assert k["offen"]["vorfragen"] == 0
    finally:
        srv.shutdown()
    # neue Sitzung, gleicher Ausgabe-Ordner: die Antwort «andere» gilt weiter, kein erneutes Fragen
    srv, url = _server_mit(tmp_path, helpers)
    try:
        req = urllib.request.Request(url + "/api/fairgate", data=fairgate, method="POST")
        with urllib.request.urlopen(req) as r:
            d = json.loads(r.read())
        assert d["vorfragen"] == [] and any(z["name"] == "Reto Brunner" for z in d["import_vorschau"])
        # zurücksetzen → Vorfrage kommt wieder
        d = _post_json(url + "/api/abgleich/entscheide", {"zuruecksetzen": True})
        assert len(d["vorfragen"]) == 1 and d["entscheide"] == {}
        # Antwort löschen (None) geht auch
        d = _post_json(url + "/api/abgleich/entscheide", {"entscheide": {"2": {"antwort": "unklar"}}})
        assert d["vorfragen"] == [] and len(d["klaerliste"]) == 1
        d = _post_json(url + "/api/abgleich/entscheide", {"entscheide": {"2": None}})
        assert len(d["vorfragen"]) == 1
    finally:
        srv.shutdown()

def test_entscheide_ungueltige_antwort_400(tmp_path):
    srv, url = _server_mit(tmp_path, [_acc_json(1, "Lina", "Brunner", "lina@example.ch", "FG-1")])
    try:
        req = urllib.request.Request(url + "/api/fairgate", method="POST", data=_fairgate_xlsx_bytes([
            ["x", 1, "lina@example.ch", "Lina", "Brunner", "", "Aktivmitglied", None, None, "2000-01-01"]]))
        urllib.request.urlopen(req).read()
        try:
            _post_json(url + "/api/abgleich/entscheide", {"entscheide": {"1": {"antwort": "vielleicht"}}})
            assert False
        except urllib.error.HTTPError as e:
            assert e.code == 400
    finally:
        srv.shutdown()
