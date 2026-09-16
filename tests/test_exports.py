import csv
import openpyxl
from cockpit.abgleich import ImportZeile, AbgleichErgebnis, HandarbeitsFall
from cockpit.exports import (IMPORT_SPALTEN, schreibe_import_xlsx,
                             schreibe_saeumigen_csv, handarbeitsliste_html,
                             schreibe_gesamtexport_xlsx)
from cockpit.model import Mitglied, Account, Typ


def _acc(vn, nn, mail):
    return Account(id=1, vorname=vn, nachname=nn, email=mail, telefon="", geburtsdatum="",
                   bemerkung="", gruppen=["Mitglied"], fg="FG-1", fg_nonstandard=False,
                   typ=Typ.MITGLIED, zielwert=2, ist_wert=0, num_ok=0, num_nok=0,
                   num_confirmed=0, num_reserved=0, num_unconfirmed=0)


def test_import_xlsx_golden(tmp_path):
    pfad = tmp_path / "import.xlsx"
    schreibe_import_xlsx([ImportZeile(vorname="Neu", nachname="Kind",
                                      email="eltern@example.ch", gruppe="Mitglied",
                                      zielwert="2", bemerkungen="FG-9")], pfad)
    ws = openpyxl.load_workbook(pfad).active
    assert [c.value for c in ws[1]] == IMPORT_SPALTEN
    zeile = [c.value for c in ws[2]]
    assert zeile[0] == "Neu" and zeile[9] == "2" and zeile[10] == "FG-9"
    assert zeile[4] == "Mitglied" and zeile[3] is None  # leere Zellen bleiben leer


def test_saeumigen_csv_folgt_der_sicht(tmp_path):
    m_halb = Mitglied(fg="FG-1", accounts=[_acc("Lina", "Brunner", "l@example.ch")], soll=2, ist=1)
    m_ok = Mitglied(fg="FG-2", accounts=[_acc("Noah", "Keller", "n@example.ch")], soll=2, ist=2)
    pfad = tmp_path / "saeumige.csv"
    n = schreibe_saeumigen_csv([m_halb, m_ok], "saison", 1, pfad)
    assert n == 1
    zeilen = list(csv.reader(pfad.open(encoding="utf-8-sig"), delimiter=";"))
    assert "Saison-Soll" in zeilen[0][0]          # Sicht in der Kopfzeile
    assert zeilen[2][0] == "Lina"                  # Kopfzeile Sicht + Spaltenkopf + Daten
    assert schreibe_saeumigen_csv([m_halb, m_ok], "halbjahr", 1, tmp_path / "h.csv") == 0


def test_handarbeitsliste_html():
    e = AbgleichErgebnis(handarbeit=[HandarbeitsFall("austritt", "Weg Gezogen", "FG-1", "deaktivieren")],
                         klaerliste=["Fall X"], zusammenfassung="1 Austritt bei 3 geprüften Mitgliedern.")
    html = handarbeitsliste_html(e)
    assert "Weg Gezogen" in html and "checkbox" in html
    assert "zuerst" in html.lower()               # Reihenfolge-Hinweis
    assert "Fall X" in html


def _zweit_acc(vn, nn, mail, fg):
    return Account(id=2, vorname=vn, nachname=nn, email=mail, telefon="", geburtsdatum="",
                   bemerkung=fg, gruppen=["Freiwillige"], fg=fg, fg_nonstandard=False,
                   typ=Typ.ZWEITACCOUNT, zielwert=0, ist_wert=1, num_ok=1, num_nok=0,
                   num_confirmed=0, num_reserved=0, num_unconfirmed=0)


def test_gesamtexport_xlsx(tmp_path):
    haupt = _acc("Lina", "Brunner", "l@example.ch")
    zweit = _zweit_acc("Rita", "Gerber", "r@example.ch", "FG-1")
    m = Mitglied(fg="FG-1", accounts=[haupt, zweit], soll=2, ist=1)
    pfad = tmp_path / "gesamt.xlsx"
    n = schreibe_gesamtexport_xlsx([m], 1, pfad)
    assert n == 1

    wb = openpyxl.load_workbook(pfad)
    assert wb.sheetnames == ["Mitglieder", "Accounts"]

    ws_m = wb["Mitglieder"]
    assert [c.value for c in ws_m[1]] == [
        "FG-Nummer", "Name", "Gruppen", "Soll", "Ist",
        "Status Saison", "Status Halbjahr", "Anzahl Accounts"]
    zeile = [c.value for c in ws_m[2]]
    assert zeile[0] == "FG-1" and zeile[1] == "Lina Brunner" and zeile[2] == "Mitglied"
    assert zeile[3] == 2 and zeile[4] == 1
    assert zeile[5] == "auf_kurs" and zeile[6] == "erfuellt" and zeile[7] == 2

    ws_a = wb["Accounts"]
    assert [c.value for c in ws_a[1]] == [
        "FG-Nummer", "Name", "Typ", "Zielwert", "Ist-Wert", "Bemerkung"]
    rows = [[c.value for c in row] for row in ws_a.iter_rows(min_row=2)]
    assert ["FG-1", "Lina Brunner", "mitglied", 2, 0, None] in rows
    assert ["FG-1", "Rita Gerber", "zweitaccount", 0, 1, "FG-1"] in rows


def test_korrektur_zeile_ueberschreibt_bemerkungen_zelle_nicht(tmp_path):
    # C1: Korrektur-Zeilen (Match über Vorname/Nachname/E-Mail-Tripel) dürfen adminRemarks im
    # Portal nie überschreiben — eine leere Bemerkungen-Zelle im Import lässt das Feld in Ruhe.
    pfad = tmp_path / "korrektur.xlsx"
    schreibe_import_xlsx([ImportZeile(vorname="Lina", nachname="Brunner",
                                      email="l@example.ch", zielwert="2", bemerkungen="")], pfad)
    ws = openpyxl.load_workbook(pfad).active
    zeile = [c.value for c in ws[2]]
    assert zeile[0] == "Lina" and zeile[9] == "2"
    assert zeile[10] is None      # Bemerkungen-Zelle bleibt leer, nicht die FG-Nummer


def test_handarbeitsliste_html_zeigt_unbekannte_kategorien():
    e = AbgleichErgebnis(
        unbekannte_kategorien=["Unbekannte Fairgate-Kategorie ‹Vereinsfremd› bei 2 Kontakten — "
                               "Regeln prüfen, diese Kontakte wurden NICHT abgeglichen."],
        zusammenfassung="Alles synchron bei 2 geprüften Mitgliedern.")
    html = handarbeitsliste_html(e)
    assert "Vereinsfremd" in html and "NICHT abgeglichen" in html


def test_handarbeitsliste_zeigt_kontaktabweichungen_als_info():
    e = AbgleichErgebnis(zusammenfassung="Alles synchron bei 1 geprüften Mitgliedern.",
                         kontakt_abweichungen=[{"name": "Lina Brunner", "fg": "FG-1", "helper_id": 1,
                                                "portal_mail": "alt@example.ch", "fairgate_mail": "neu@example.ch"}])
    html = handarbeitsliste_html(e)
    assert "Kontaktdaten weichen ab" in html and "alt@example.ch" in html and "neu@example.ch" in html
    assert "Nichts zu tun" in html          # Info-Liste ist keine Handarbeit

def test_kontaktabweichungen_csv(tmp_path):
    from cockpit.exports import schreibe_kontaktabweichungen_csv
    pfad = tmp_path / "k.csv"
    n = schreibe_kontaktabweichungen_csv([{"name": "Lina Brunner", "fg": "FG-1", "helper_id": 1,
                                           "portal_mail": "alt@example.ch", "fairgate_mail": "neu@example.ch"}], pfad)
    assert n == 1
    zeilen = list(csv.reader(pfad.open(encoding="utf-8-sig"), delimiter=";"))
    assert zeilen[0][:4] == ["Name", "FG-Nummer", "E-Mail Portal", "E-Mail Fairgate"] and zeilen[1][2] == "alt@example.ch"


def test_handarbeitsliste_rendert_strukturierte_klaerfaelle():
    e = AbgleichErgebnis(zusammenfassung="Alles synchron bei 1 geprüften Mitgliedern.",
                         klaerliste=[{"titel": "Petra: gleiche E-Mail wie Jan (FG-5)", "wo": "Portal",
                                      "fakten": ["Fakt eins"], "optionen": ["Option A", "Option B"], "helper_id": 2, "fg": "FG-5"}])
    html = handarbeitsliste_html(e)
    assert "Klärfälle" in html and "Fakt eins" in html and "→ Option B" in html and "(Portal)" in html


def test_import_begruendung_html_und_kategorien():
    from cockpit.abgleich import ImportZeile
    from cockpit.exports import import_begruendung_html, import_zeilen_mit_grund, grund_kategorie
    e = AbgleichErgebnis(
        neueintritte=[ImportZeile("Neu", "Kind", "n@example.ch", gruppe="Mitglied", zielwert="2", bemerkungen="FG-9",
                                  grund="Neueintritt: in Fairgate, noch kein Portal-Account")],
        korrekturen=[ImportZeile("Karl", "Ohne", "k@example.ch", zielwert="0", gruppe="Freiwillige",
                                 grund="Kein Mitglied in Fairgate (weder FG, E-Mail noch Name): Freiwillige(r), Zielwert 0"),
                     ImportZeile("Lina", "Brunner", "l@example.ch", zielwert="2", grund="Mitglied: Zielwert nachführen")])
    z = import_zeilen_mit_grund(e)
    assert [x["zeile"] for x in z] == [2, 3, 4]
    assert [x["kategorie"] for x in z] == ["Neueintritt", "Kein Mitglied in Fairgate → Freiwillige", "Mitglied aktualisieren"]
    assert z[0]["aenderungen"].startswith("E-Mail n@example.ch") and "Bemerkung FG-9" in z[0]["aenderungen"]
    assert grund_kategorie("Zweitaccount von FG-3 («X», gleiche E-Mail): …") == "Zweitaccount"
    assert grund_kategorie("Kein Mitglied in Fairgate (…): Freiwillige(r), Zielwert 0 — andere Person als das namensgleiche Mitglied (deine Antwort)") == "Kein Mitglied in Fairgate → Freiwillige"
    html = import_begruendung_html(e, "import-2026-09-10.xlsx")
    assert "import-2026-09-10.xlsx" in html and "Karl Ohne" in html and "<td class='n'>3</td>" in html
    assert "Kein Mitglied in Fairgate → Freiwillige: 1" in html
