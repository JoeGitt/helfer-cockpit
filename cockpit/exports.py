"""Ausgabedateien: Import-Excel (Portal-Format), Säumigen-CSV, Handarbeitsliste."""
import csv
import html as html_mod
from pathlib import Path
import openpyxl
from .model import status, Typ

IMPORT_SPALTEN = ["Vorname", "Nachname", "E-Mail", "Telefon", "Gruppe",
                  "zusätzliche E-Mail (1)", "zusätzliche E-Mail (2)", "Geburtsdatum",
                  "Geschlecht", "Individueller Zielwert", "Bemerkungen"]

_SICHT_NAME = {"saison": "Saison-Soll", "halbjahr": "Halbjahresziel"}


def schreibe_import_xlsx(zeilen, pfad):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Import"
    ws.append(IMPORT_SPALTEN)
    for z in zeilen:
        ws.append([z.vorname or None, z.nachname or None, z.email or None, z.telefon or None,
                   z.gruppe or None, z.zusatz_email1 or None, z.zusatz_email2 or None,
                   z.geburtsdatum or None, z.geschlecht or None, z.zielwert or None,
                   z.bemerkungen or None])
    wb.save(pfad)


MITGLIEDER_SPALTEN = ["FG-Nummer", "Name", "Gruppen", "Soll", "Ist",
                      "Status Saison", "Status Halbjahr", "Anzahl Accounts", "Familie", "Familie Soll", "Familie Ist"]
ACCOUNTS_SPALTEN = ["FG-Nummer", "Name", "Typ", "Zielwert", "Ist-Wert", "Bemerkung"]
ALLE_SPALTEN = ["ID", "Name", "Typ", "Gruppen", "FG-Nummer", "Geleistet (OK)",
                "Nicht erschienen (NOK)", "Zugesagt", "Reserviert", "Ist-Wert", "Zielwert",
                "Bemerkung"]


def schreibe_gesamtexport_xlsx(mitglieder, halbjahresziel, pfad, accounts=None):
    """Excel-Gesamtexport (Spez. 6.1): Blatt «Mitglieder», Blatt «Accounts» (Accounts der
    Mitglieder) und — wenn übergeben — Blatt «Alle Helfenden» mit jedem Portal-Account."""
    wb = openpyxl.Workbook()
    ws_m = wb.active
    ws_m.title = "Mitglieder"
    ws_m.append(MITGLIEDER_SPALTEN)
    ws_a = wb.create_sheet("Accounts")
    ws_a.append(ACCOUNTS_SPALTEN)
    for m in mitglieder:
        a0 = m.mitglieds_account
        ws_m.append([m.fg, a0.anzeigename, ", ".join(a0.gruppen), m.soll, m.ist,
                    status(m, "saison", halbjahresziel), status(m, "halbjahr", halbjahresziel),
                    len(m.accounts), m.familie.name if m.familie else "",
                    m.familie.soll if m.familie else None, m.familie.ist if m.familie else None])
        for a in m.accounts:
            ws_a.append([m.fg, a.anzeigename, a.typ.value, a.zielwert, a.ist_wert, a.bemerkung])
    if accounts is not None:
        ws_h = wb.create_sheet("Alle Helfenden")
        ws_h.append(ALLE_SPALTEN)
        for a in sorted(accounts, key=lambda x: (x.nachname, x.vorname)):
            ws_h.append([a.id, a.anzeigename, a.typ.value, ", ".join(a.gruppen), a.fg or "",
                         a.num_ok, a.num_nok, a.num_confirmed, a.num_reserved,
                         a.ist_wert, a.zielwert, a.bemerkung])
    wb.save(pfad)
    return len(mitglieder)


def schreibe_saeumigen_csv(mitglieder, sicht, halbjahresziel, pfad):
    """Eine Zeile pro säumigem Mitglied; eine Familie (Topf) erscheint als eine Zeile mit allen
    Kindern und der E-Mail des Elternaccounts, damit nur eine Erinnerung an die Familie geht."""
    saeumige, familien_gesehen = [], set()
    for m in mitglieder:
        if status(m, sicht, halbjahresziel) == "erfuellt":
            continue
        if m.familie:
            if id(m.familie) in familien_gesehen:
                continue
            familien_gesehen.add(id(m.familie))
        saeumige.append(m)
    with Path(pfad).open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow([f"Säumigen-Liste — Sicht: {_SICHT_NAME[sicht]}"])
        w.writerow(["Vorname", "Nachname", "E-Mail", "FG-Nummer", "Soll", "Ist"])
        for m in saeumige:
            a = m.mitglieds_account
            if m.familie:
                fam = m.familie
                kinder = [x for x in fam.accounts if x.typ == Typ.MITGLIED and x.fg in fam.fgs]
                eltern = [x for x in fam.accounts if x.typ != Typ.MITGLIED]
                n = len(fam.fgs)
                soll = fam.soll if sicht == "saison" else halbjahresziel * n
                w.writerow([" + ".join(x.vorname for x in kinder), fam.name.replace("Familie ", ""),
                            (eltern[0].email if eltern else a.email), ", ".join(fam.fgs),
                            f"{soll:g}", f"{fam.ist:g}"])
            else:
                w.writerow([a.vorname, a.nachname, a.email, m.fg, f"{m.soll:g}", f"{m.ist:g}"])
    return len(saeumige)


# Kurzlabel je Begründung — für Filter-Chips und die Spalte «Warum»
_GRUND_KATEGORIEN = (("Neueintritt", "Neueintritt"), ("Zweitaccount", "Zweitaccount"),
                     ("Ersetzt den bisherigen", "Ersatz-Account"), ("Bisheriger Account", "Ersatz-Account"),
                     ("Ist Mitglied", "FG-Nummer nachtragen"),
                     ("Kein Mitglied", "Kein Mitglied in Fairgate → Freiwillige"),
                     ("Unbekannt", "Unbekannt mit Einsätzen → Freiwillige"),
                     ("ohne Helferpflicht", "Keine Helferpflicht → Zielwert 0"),
                     ("ohne FG-Nummer", "FG-Nummer nachtragen"), ("Ist das Mitglied", "FG-Nummer nachtragen"))


def grund_kategorie(grund):
    for schluessel, label in _GRUND_KATEGORIEN:
        if schluessel in (grund or ""):
            return label
    return "Mitglied aktualisieren"


def import_zeilen_mit_grund(e):
    """Jede Import-Zeile mit Excel-Zeilennummer, Art, Kurzgrund, Begründung und den Feldern, die
    geschrieben werden. Reihenfolge = Reihenfolge in der Datei (Neueintritte, dann Korrekturen)."""
    def felder(z, neu):
        teile = []
        if neu and z.email: teile.append(f"E-Mail {z.email}")
        if z.gruppe: teile.append(f"Gruppen → {z.gruppe}")
        if z.zielwert: teile.append(f"Zielwert {z.zielwert}")
        if z.telefon: teile.append(f"Telefon {z.telefon}")
        if z.bemerkungen: teile.append(f"Bemerkung {z.bemerkungen}")
        return ", ".join(teile)
    zeilen = []
    for i, (z, neu) in enumerate([(z, True) for z in e.neueintritte] + [(z, False) for z in e.korrekturen]):
        zeilen.append({"zeile": i + 2, "name": f"{z.vorname} {z.nachname}".strip(), "email": z.email,
                       "art": "Neueintritt" if neu else "Korrektur", "kategorie": grund_kategorie(z.grund),
                       "grund": z.grund, "aenderungen": felder(z, neu)})
    return zeilen


def import_begruendung_html(e, import_dateiname=""):
    """Druckbare Begleitdatei zur Import-Datei: warum jede Zeile drin ist. Wird nicht importiert."""
    zeilen = import_zeilen_mit_grund(e)
    esc = html_mod.escape
    zaehler = {}
    for z in zeilen:
        zaehler[z["kategorie"]] = zaehler.get(z["kategorie"], 0) + 1
    teile = ["<meta charset='utf-8'><title>Begründung zur Import-Datei</title>",
             "<style>body{font:14px/1.5 sans-serif;max-width:1100px;margin:2rem auto;padding:0 1rem}"
             "table{border-collapse:collapse;width:100%}th,td{text-align:left;vertical-align:top;"
             "padding:6px 8px;border-bottom:1px solid #ddd}th{font-size:12px;text-transform:uppercase;"
             "letter-spacing:.05em;color:#555}td.n{text-align:right;font-variant-numeric:tabular-nums}"
             "tr{break-inside:avoid}</style>",
             f"<h1>Begründung zur Import-Datei</h1><p><b>{esc(import_dateiname)}</b> · {len(zeilen)} Zeilen · "
             "Diese Datei ist nur zum Nachlesen und wird nicht ins Portal importiert. Die Zeilennummer "
             "entspricht der Zeile in Excel (Kopfzeile = 1).</p>",
             "<p>" + " · ".join(f"{esc(k)}: {n}" for k, n in sorted(zaehler.items(), key=lambda kv: -kv[1])) + "</p>",
             "<table><thead><tr><th>Zeile</th><th>Person</th><th>E-Mail</th><th>Art</th><th>Warum</th>"
             "<th>Begründung</th><th>Was geschrieben wird</th></tr></thead><tbody>"]
    for z in zeilen:
        teile.append(f"<tr><td class='n'>{z['zeile']}</td><td>{esc(z['name'])}</td><td>{esc(z['email'])}</td>"
                     f"<td>{esc(z['art'])}</td><td>{esc(z['kategorie'])}</td><td>{esc(z['grund'])}</td>"
                     f"<td>{esc(z['aenderungen'])}</td></tr>")
    teile.append("</tbody></table>")
    if not zeilen:
        teile.append("<p>Keine Import-Zeilen — dieses Mal ist kein Import nötig.</p>")
    return "\n".join(teile)


def handarbeitsliste_html(e):
    def punkt(text):
        return ('<li><label><input type="checkbox"> ' + html_mod.escape(text) + "</label></li>")
    schluessel = [h for h in e.handarbeit if h.art == "schluessel"]
    austritte = [h for h in e.handarbeit if h.art == "austritt"]
    teile = ["<meta charset='utf-8'><title>Handarbeits-Liste</title>",
             "<style>body{font:15px/1.6 sans-serif;max-width:720px;margin:2rem auto;padding:0 1rem}"
             "li{margin:.4rem 0}@media print{input{-webkit-print-color-adjust:exact}}</style>",
             f"<h1>Handarbeits-Liste</h1><p><b>{html_mod.escape(e.zusammenfassung)}</b></p>",
             "<p><b>Reihenfolge:</b> zuerst diese Liste im Portal abarbeiten "
             "(Schlüssel-Änderungen vor allem!), erst danach die Import-Datei hochladen — "
             "sonst entstehen Duplikate.</p>"]
    if schluessel:
        teile.append("<h2>1 · Schlüssel-Änderungen (zuerst!)</h2><ul>")
        teile += [punkt(f"{h.name} ({h.fg}): {h.detail}") for h in schluessel]
        teile.append("</ul>")
    if austritte:
        teile.append("<h2>2 · Austritte im Portal löschen</h2>"
                     "<p>Das Portal kennt kein Deaktivieren: Helfende → Person → «Helfer:in löschen» "
                     "(unwiderruflich; vergangene Einsätze verschwinden aus der Statistik).</p><ul>")
        teile += [punkt(f"{h.name} ({h.fg}): {h.detail}") for h in austritte]
        teile.append("</ul>")
    if e.duplikat_warnungen:
        teile.append("<h2>Duplikat-Warnungen (nicht importiert)</h2><ul>")
        teile += [punkt(t) for t in e.duplikat_warnungen]
        teile.append("</ul>")
    if e.klaerliste:
        teile.append("<h2>Klärfälle — hier musst du entscheiden</h2><ul>")
        for kf in e.klaerliste:
            if isinstance(kf, dict):
                teile.append("<li><label><input type=\"checkbox\"> <b>" + html_mod.escape(kf["titel"])
                             + f"</b> <i>({html_mod.escape(kf.get('wo', 'Portal'))})</i></label><ul>"
                             + "".join(f"<li>{html_mod.escape(f)}</li>" for f in kf.get("fakten", []))
                             + "</ul><ul>" + "".join(f"<li>→ {html_mod.escape(o)}</li>" for o in kf.get("optionen", []))
                             + "</ul></li>")
            else:
                teile.append(punkt(str(kf)))
        teile.append("</ul>")
    if e.unbekannte_kategorien:
        teile.append("<h2>Unbekannte Fairgate-Kategorien (Regeln prüfen)</h2><ul>")
        teile += [punkt(t) for t in e.unbekannte_kategorien]
        teile.append("</ul>")
    if not (schluessel or austritte or e.duplikat_warnungen or e.klaerliste
            or e.unbekannte_kategorien):
        teile.append("<p>Nichts zu tun — alles synchron. 🎉</p>")
    hinweise = getattr(e, "hinweise", [])
    if hinweise:
        teile.append("<h2>Hinweise — der Import erledigt sie, optional prüfen</h2><ul>")
        for h in hinweise:
            if isinstance(h, dict):
                teile.append("<li><b>" + html_mod.escape(h["titel"]) + "</b><ul>"
                             + "".join(f"<li>{html_mod.escape(f)}</li>" for f in h.get("fakten", []))
                             + "</ul><ul>" + "".join(f"<li>→ {html_mod.escape(o)}</li>" for o in h.get("optionen", []))
                             + "</ul></li>")
            else:
                teile.append(f"<li>{html_mod.escape(str(h))}</li>")
        teile.append("</ul>")
    abw = getattr(e, "kontakt_abweichungen", [])
    if abw:
        teile.append("<h2>Kontaktdaten weichen ab (Info — keine Handarbeit nötig)</h2>"
                     "<p>Die Portal-Adresse ist die vom Mitglied selbst gewählte Login-Adresse. "
                     "Falls Fairgate veraltet ist, dort nachführen; im Portal ist nichts zu tun.</p><ul>")
        teile += [f"<li>{html_mod.escape(a['name'])} ({html_mod.escape(a['fg'])}): Portal "
                  f"{html_mod.escape(a['portal_mail'])} · Fairgate {html_mod.escape(a['fairgate_mail'])}</li>"
                  for a in abw]
        teile.append("</ul>")
    return "\n".join(teile)


def schreibe_kontaktabweichungen_csv(abweichungen, pfad):
    """Info-Liste «E-Mail Portal ≠ Fairgate» als CSV — zum Nachführen in Fairgate."""
    with Path(pfad).open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["Name", "FG-Nummer", "E-Mail Portal", "E-Mail Fairgate"])
        for a in abweichungen:
            w.writerow([a["name"], a["fg"], a["portal_mail"], a["fairgate_mail"]])
    return len(abweichungen)
